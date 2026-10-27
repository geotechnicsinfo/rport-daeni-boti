#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Construiește data.json (+ pozele optimizate) pentru aplicația de raportare a investigațiilor de teren –
Parc eolian Dăeni, jud. Tulcea.

Surse:
  - De la ei fise/Fise Daeni +rezultate lab (WT1-WT56).xlsx – fișele formatate ale forajelor (GEO-TECH SRL), o foaie per
    foraj („1+r”, „2”, … „35-nou”, „39b”, „55+r”): antet cu Stereo 70 x/y (x = est, y = nord) și cota RNMN, stratificație,
    probe (US/DS) cu rezultatele de laborator pe același rând, apă, SPT (coloanele AF–AH)
  - Turbine_Daeni_nov_2025_wt 11.11.2025.kmz – planul turbinelor WT1–WT56 (revizuit 11.11.2025)
  - cpt.json (generat de build_cpt.py) – CPT-ul turbinei (același punct) și sumarul CPT
  - scripts/meteo.json (Open-Meteo archive, opțional); Poze/<ID> (opțional) – poze WhatsApp din campania CPT;
    folderul „ZONA WT41 WT42 CEA MAI INALTA” -> galeria de zonă

Rulare (din folderul aplicației):
  python3 scripts/build_data.py --fise "../De la ei fise/Fise Daeni +rezultate lab (WT1-WT56).xlsx" \
     --kmz "../Turbine_Daeni_nov_2025_wt 11.11.2025.kmz" --cpt cpt.json --meteo scripts/meteo.json --poze "../Poze" --out .

Conversia Stereo 70 → WGS84 folosește pyproj (EPSG:3844 → EPSG:4326); fără pyproj se citește scripts/foraje_wgs84.json
(cache scris la ultima rulare cu pyproj).
"""
import argparse, html, json, math, os, re, datetime, zipfile
import openpyxl

MAX_W, THUMB_W, QUALITY = 1200, 420, 68
ZONE_NUME = {'N': 'Zona N – WT1–WT23', 'S': 'Zona S – WT24–WT56'}
ADANCIME_PROIECTATA = 30.0
POZE_ALIAS = {'WT015': 'WT15', 'WT7b': 'WT7', 'WT9b': 'WT9', 'WT17b': 'WT17', 'WT35nou': 'WT35'}
POZE_ZONA = 'ZONA WT41 WT42 CEA MAI INALTA'
LAB_KEYS = ['A', 'P', 'N', 'G', 'B', 'w', 'wL', 'wP', 'IP', 'IC', 'gama', 'gama_d', 'n', 'e', 'Sr', 'M23', 'ep2', 'im3', 'phi', 'c']


def zona_of(iid):
    m = re.match(r'WT(\d+)', iid); return 'N' if m and int(m.group(1)) <= 23 else 'S'


def fnum(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(',', '.')
    if s in ('', '--', '-', 'None', '—'): return None
    try: return float(s)
    except ValueError: return None


def rdepth(d):
    d = round(float(d), 3)
    if abs(d - round(d)) < 0.015: return float(round(d))
    return round(d, 2)


def clean_desc(s):
    s = re.sub(r'\s+', ' ', str(s)).strip(' ,')
    s = s.replace('↗', '').replace('Loes ', 'Loess ')
    s = re.sub(r'\s*/\s*', ' / ', s)
    s = re.sub(r'\s+,', ',', s)
    return s.strip()


def merge_desc(descs):
    uniq = []
    for d in descs:
        if d not in uniq: uniq.append(d)
    if len(uniq) == 1: return uniq[0]
    base = uniq[0].rsplit(',', 1)
    if len(base) == 2 and all(u.rsplit(',', 1)[0] == base[0] for u in uniq):
        states = []
        for u in uniq:
            for st in u.rsplit(',', 1)[1].split('/'):
                st = st.strip()
                if st and st not in states: states.append(st)
        return f"{base[0]}, {' / '.join(states)}"
    return ' | '.join(uniq)


def parse_fisa(ws):
    rows = list(ws.iter_rows(values_only=True))
    iid = x = y = z0 = None; water = []; descs = []; bounds = []; bottom = 0.0; spt = []; probe = []
    for r in rows:
        r = list(r) + [None] * (40 - len(r))
        if r[15] == 'FIŞA FORAJULUI GEOTEHNIC' and iid is None: iid = str(r[16]).strip()
        if r[7] == 'Coord. Stereo 70' and x is None: x, y, z0 = fnum(r[9]), fnum(r[15]), fnum(r[2])
        d = r[3]
        if isinstance(d, (int, float)) and r[2] != '(m)':
            d = rdepth(d)
            bottom = max(bottom, d)
            if isinstance(r[4], (int, float)) and r[4] > 0 and (not bounds or d > bounds[-1] + 1e-6): bounds.append(d)
            if isinstance(r[7], str) and r[7].strip() and r[7].strip() not in ('Descrierea stratului', 'Verificat:'): descs.append((d, r[7]))
            if isinstance(r[5], str) and r[5].strip() and r[5].strip() != '(m)' and not r[5].startswith(('NA-', 'INF')):
                water.append((d, r[5].strip()))
            if isinstance(r[8], str) and re.match(r'(US|DS)\d', r[8].strip()):
                iv = str(r[9] or '').strip().replace(',', '.')
                m = re.match(r'([\d.]+)\s*[-–_]\s*([\d.]+)', iv)
                de, la = (fnum(m.group(1)), fnum(m.group(2))) if m else (None, fnum(iv))
                p = {'proba': r[8].strip(), 'netulburata': r[8].strip().startswith('US'), 'tip': 'US' if r[8].strip().startswith('US') else 'DS',
                     'interval': iv.replace('.', ','), 'de': de, 'adancime': la if la is not None else d, 'denumire': clean_desc(r[7]) if r[7] else None}
                for k, v in zip(LAB_KEYS, r[10:30]): p[k] = fnum(v)
                if p['IC'] is None and p['w'] is not None and p['wL'] is not None and p['IP']: p['IC'] = round((p['wL'] - p['w']) / p['IP'], 2)
                probe.append(p)
        if isinstance(r[31], str) and r[31].strip() and r[33] is not None:
            m = re.match(r'([\d,.]+)\s*[_\-–]\s*([\d,.]+)', r[31].strip())
            if m: spt.append({'de': fnum(m.group(1)), 'la': fnum(m.group(2)), 'valori': str(r[32]).replace('_', '-'), 'N': fnum(r[33])})
    if bottom > (bounds[-1] if bounds else 0) + 1e-6: bounds.append(bottom)
    layers = []; top = 0.0
    for bot in bounds:
        ds = [clean_desc(t) for d, t in descs if top < d <= bot + 1e-6]
        if not ds and layers: ds = [layers[-1]['descriere']]
        desc = merge_desc(ds) if ds else '—'
        layers.append({'de': top, 'la': bot, 'descriere': desc}); top = bot
    for k, L in enumerate(layers): L['nr'] = k + 1
    # SPT: dedublare (rândurile de continuare repetă uneori ultima încercare)
    seen = set(); spt2 = []
    for s in spt:
        key = (s['de'], s['la'])
        if key in seen: continue
        seen.add(key); spt2.append(s)
    for s in spt2: s['N'] = round(s['N']) if s['N'] is not None else None
    return {'id': iid, 'x': x, 'y': y, 'cota': z0, 'water': water, 'straturi': layers, 'adancime': bottom, 'spt': sorted(spt2, key=lambda s: s['de']),
            'probe': sorted(probe, key=lambda p: (p['adancime'] or 0, p['proba']))}


def read_kmz(path):
    if path.lower().endswith('.kmz'):
        with zipfile.ZipFile(path) as zf:
            s = zf.read(next(n for n in zf.namelist() if n.lower().endswith('.kml'))).decode('utf-8')
    else:
        s = open(path, encoding='utf-8').read()
    out = {}
    for p in re.findall(r'<Placemark>(.*?)</Placemark>', s, re.S):
        n = html.unescape(re.search(r'<name>(.*?)</name>', p, re.S).group(1).strip())
        c = re.search(r'<coordinates>(.*?)</coordinates>', p, re.S)
        if not c: continue
        x = c.group(1).split()[0].split(',')
        out[n] = {'lat': round(float(x[1]), 7), 'lon': round(float(x[0]), 7)}
    return out


def dist(a, b, c, d):
    if None in (a, b, c, d): return None
    R = 6371000; t = math.pi / 180; dl = (c - a) * t; dn = (d - b) * t
    h = math.sin(dl / 2) ** 2 + math.cos(a * t) * math.cos(c * t) * math.sin(dn / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))


def process_photos(folder, out_dir, rel, start=0):
    from PIL import Image, ImageOps
    os.makedirs(out_dir, exist_ok=True)
    files = sorted((f for f in os.listdir(folder) if re.search(r'\.jpe?g$|\.png$', f, re.I)), key=str.lower)
    fise, pics, n, nf = [], [], start, 0
    for f in files:
        is_fisa = 'fisa' in f.lower()
        src = os.path.join(folder, f)
        try: im = ImageOps.exif_transpose(Image.open(src)).convert('RGB')
        except Exception as e: print('  ! poza ignorată', src, e); continue
        if is_fisa: nf += 1; name = f'fisa{nf}'
        else: n += 1; name = f'{n:02d}'
        for w, suffix in ((MAX_W, ''), (THUMB_W, 't_')):
            dst = os.path.join(out_dir, f'{suffix}{name}.jpg')
            if os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src): continue
            img = im.copy(); img.thumbnail((w, w)); img.save(dst, 'JPEG', quality=QUALITY, optimize=True, progressive=True)
        entry = {'src': f'{rel}/{name}.jpg', 'thumb': f'{rel}/t_{name}.jpg', 'original': f, 'w': im.width, 'h': im.height}
        m = re.search(r'(\d{4}-\d{2}-\d{2}) at (\d{1,2})\.(\d{2})\.(\d{2})(?:\s*([AP]M))?', f)
        if m:
            hh = int(m.group(2))
            if m.group(5) == 'PM' and hh < 12: hh += 12
            if m.group(5) == 'AM' and hh == 12: hh = 0
            entry['timestamp'] = f'{m.group(1)} {hh:02d}:{m.group(3)}'; entry['data'] = m.group(1)
        (fise if is_fisa else pics).append(entry)
    return fise, pics


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--fise', required=True); ap.add_argument('--kmz', required=True)
    ap.add_argument('--cpt', default='cpt.json'); ap.add_argument('--meteo', default=None)
    ap.add_argument('--poze', default=None); ap.add_argument('--out', required=True)
    a = ap.parse_args()
    here = os.path.dirname(os.path.abspath(__file__))

    plan = read_kmz(a.kmz)
    cptj = json.load(open(a.cpt, encoding='utf-8')) if os.path.exists(a.cpt) else {'meta': {}, 'teste': []}
    tests = cptj['teste']

    cache_path = os.path.join(here, 'foraje_wgs84.json'); cache = {}
    try:
        from pyproj import Transformer
        tr = Transformer.from_crs('EPSG:3844', 'EPSG:4326', always_xy=True)
        def to_wgs(E, N):
            lon, lat = tr.transform(E, N); return round(lat, 7), round(lon, 7)
    except Exception:
        tr = None
        cache = json.load(open(cache_path, encoding='utf-8')) if os.path.exists(cache_path) else {}
        def to_wgs(E, N):
            v = cache.get(f'{E:.2f},{N:.2f}'); return (v[0], v[1]) if v else (None, None)

    fise = {}
    wb = openpyxl.load_workbook(a.fise, data_only=True)
    for name in wb.sheetnames:
        f = parse_fisa(wb[name])
        if not f['id']: print('! foaie fără antet', name); continue
        fise[f['id']] = dict(f, foaie=name, fisier=os.path.basename(a.fise))
    print('fișe:', len(fise), 'probe:', sum(len(v['probe']) for v in fise.values()), 'SPT:', sum(len(v['spt']) for v in fise.values()), 'CPT:', len(tests))

    # poze: foldere pe turbină (alias-uri) + galeria de zonă
    poze_dir = {}; zona_gal = []
    if a.poze and os.path.isdir(a.poze):
        for d in sorted(os.listdir(a.poze)):
            p = os.path.join(a.poze, d)
            if not os.path.isdir(p): continue
            if d == POZE_ZONA:
                _, zona_gal = process_photos(p, os.path.join(a.out, 'photos', 'zona'), 'photos/zona'); continue
            key = POZE_ALIAS.get(d, d)
            if not re.match(r'WT\d+$', key): print('  poze ignorate (nu sunt ale unei turbine):', d); continue
            poze_dir.setdefault(key, []).append((d, p))

    order = sorted(plan, key=lambda b: int(re.sub(r'\D', '', b)))
    investigatii = []; wgs_cache = {}
    for iid in order:
        f = fise.get(iid)
        if not f: print('! lipsă fișă', iid); continue
        lat, lon = to_wgs(f['x'], f['y']); wgs_cache[f"{f['x']:.2f},{f['y']:.2f}"] = [lat, lon]
        pl = plan.get(iid)
        my_cpt = [t for t in tests if t['turbina'] == iid]
        note = ['Data execuției nu este consemnată în fișa forajului; CPT-ul turbinei are data în antetul fișierului Geo-Explorer.']
        dcpt = None
        if my_cpt and lat is not None:
            cl = [t for t in my_cpt if t['lat'] is not None and t['nr'] == 1]
            if cl:
                dcpt = dist(lat, lon, cl[0]['lat'], cl[0]['lon'])
        if pl and lat is not None:
            dp = dist(lat, lon, pl['lat'], pl['lon'])
            if dp > 30: note.append(f'Coordonatele din fișă sunt la ~{dp:.0f} m de punctul din plan (KMZ 11.11.2025).')
        if iid == 'WT35':
            note.append('Fișa „35-nou” păstrează în antet coordonatele poziției inițiale; turbina are două CPT-uri (WT35 pe poziția din plan, WT35nou pe locația nouă, 2,5 km NNE).')
        if f['adancime'] < ADANCIME_PROIECTATA - 0.01:
            note.append(f"Foraj oprit la {f['adancime']:.2f} m (din {ADANCIME_PROIECTATA:.0f} m proiectați) – șist verde.".replace('.', ',', 1).replace('.00', ',00'))
        allw = ' '.join(w for _, w in f['water'])
        m = re.search(r'NA\s*-?\s*([\d.,]+)', allw); NH = fnum(m.group(1)) if m else None
        m2 = re.search(r'Apă\s*≈\s*([\d.,]+)', allw)
        if NH is None and m2: NH = fnum(m2.group(1))
        m3 = re.search(r'infiltra\w*\s*(?:apă)?\s*-?\s*([\d.,]+)', allw, re.I); IN = fnum(m3.group(1)) if m3 else None
        iv = re.search(r'([\d]+[.,]\d)\s*-\s*([\d]+[.,]\d)', allw)
        ivs = f'{iv.group(1)}–{iv.group(2)} m'.replace('.', ',') if iv else None
        if NH is not None: apa = f'NA {ivs}' if ivs else f'NA ≈ {NH:.1f} m'.replace('.', ',')
        elif IN is not None: apa = f'infiltrații de apă {ivs}' if ivs else f'infiltrații de apă la {IN:.1f} m'.replace('.', ',')
        elif 'infiltra' in allw.lower(): apa = 'infiltrații de apă (fără adâncime consemnată)'
        else: apa = 'fără apă'
        poze, fise_img = [], []
        for d, p in poze_dir.get(iid, []):
            print('poze', iid, d)
            fi, pz = process_photos(p, os.path.join(a.out, 'photos', iid), f'photos/{iid}', start=len(poze))
            for e in pz: e['folder'] = d
            poze += pz; fise_img += fi
        probe = f['probe']
        investigatii.append({
            'id': iid, 'id_fisa': f['foaie'], 'tip': 'foraj', 'zona': zona_of(iid), 'zona_nume': ZONE_NUME[zona_of(iid)],
            'turbina': iid, 'grup': iid, 'amplasare': f'Turbina {iid}',
            'categorie_kml': 'Foraj geotehnic + SPT – fundație turbină',
            'eticheta': f'Foraj geotehnic {iid}',
            'data': None, 'nota_data': ' '.join(note),
            'adancime': f['adancime'], 'adancime_proiectata': ADANCIME_PROIECTATA,
            'apa': apa, 'NH': NH, 'NH_24h': None, 'IN': IN, 'note_apa': [w for _, w in f['water']],
            'cota': f['cota'], 'stereo70': {'E': f['x'], 'N': f['y']}, 'lat': lat, 'lon': lon,
            'coord_sursa': 'fișa forajului (Stereo 70 → WGS84)', 'gps_dms': [],
            'proiectat': {'id': iid, 'lat': pl['lat'], 'lon': pl['lon'], 'obs': 'plan turbine 11.11.2025', 'adancime': '30 m'} if pl else None,
            'pietris': None,
            'straturi': f['straturi'],
            'probe': [{'nr': k + 1, 'cod': p['proba'], 'adancime': p['adancime'], 'tip': 'probă netulburată (US)' if p['netulburata'] else 'probă tulburată (DS)', 'observatii': p['denumire'] or ''} for k, p in enumerate(probe)],
            'lab': probe, 'spt': f['spt'],
            'cpt_ids': [t['id'] for t in my_cpt], 'dist_cpt': round(dcpt) if dcpt is not None else None,
            'fisa_foto': fise_img[0] if fise_img else None, 'fise_foto': fise_img, 'poze': poze,
        })
    if tr: json.dump(wgs_cache, open(cache_path, 'w', encoding='utf-8'), indent=0)

    meteo = {}
    if a.meteo and os.path.exists(a.meteo):
        mj = json.load(open(a.meteo, encoding='utf-8')); d = mj.get('daily', {})
        for i, t in enumerate(d.get('time', [])): meteo[t] = {k: d[k][i] for k in d if k != 'time'}
        meteo['_meta'] = {'sursa': 'Open-Meteo (ERA5 / archive API)', 'lat': mj.get('latitude'), 'lon': mj.get('longitude'), 'elevatie_m': mj.get('elevation')}

    stadiu = []
    exec_cpt_turb = {t['turbina'] for t in tests}
    for z in ('N', 'S'):
        inv = [i for i in investigatii if i['zona'] == z]; tz = [t for t in tests if t['zona'] == z]
        bp = [b for b in order if zona_of(b) == z]
        stadiu.append({'zona': z, 'nume': ZONE_NUME[z],
                       'foraje': {'executate': len(inv), 'proiectate': len(bp), 'neexecutate': [b for b in bp if b not in {i['id'] for i in inv}]},
                       'cptu': {'executate': len([b for b in bp if b in exec_cpt_turb]), 'proiectate': len(bp), 'neexecutate': [b for b in bp if b not in exec_cpt_turb], 'teste': len(tz)},
                       'm_foraj': round(sum(i['adancime'] for i in inv), 1), 'm_cpt': round(sum(t['adancime'] for t in tz), 2),
                       'probe': sum(len(i['lab']) for i in inv), 'probe_us': sum(1 for i in inv for p in i['lab'] if p['netulburata']),
                       'spt': sum(len(i['spt']) for i in inv)})

    light = [{k: t[k] for k in ('id', 'tip', 'grup', 'nr', 'turbina', 'zona', 'data', 'adancime', 'lat', 'lon', 'stereo70', 'coord_sursa', 'nota', 'gwl', 'straturi_sbt', 'manual')}
             | {'qc_med': t['sumar']['qc_med'], 'qc_max': t['sumar']['qc_max'], 'u2_max': t['sumar']['u2_max'],
                'benzi_1m': [{'de': b['de'], 'la': b['la'], 'qc_med': b['qc_med'], 'qc_max': b['qc_max']} for b in t['benzi_1m']]} for t in tests]
    M = cptj.get('meta', {})
    out = {
        'proiect': {
            'titlu': 'Investigații geotehnice – Parc eolian Dăeni',
            'proiect': 'Construire Parc eolian Dăeni, drumuri noi de acces, stații de transformare, rețele electrice – UAT Dăeni, jud. Tulcea',
            'beneficiar': 'SC OXIGEN DELTA SRL', 'antreprenor': 'ARVADA SRL',
            'elaborator_studiu': 'GEO-TECH SRL Gheorgheni (foraje, laborator, studiu geotehnic)',
            'beneficiar_lucrare': 'Fundațiile turbinelor eoliene WT1–WT56',
            'amplasament': 'UAT Dăeni, jud. Tulcea',
            'executant': 'Foraje + laborator: GEO-TECH SRL · CPT: BBCGEOTEHNIC SRL',
            'tip_lucrari': 'Foraje geotehnice cu prelevare de probe și SPT (30 m, în centrul fiecărei turbine), penetrări statice CPT (con electric, câte una per turbină, ≤ 25,5 m) și încercări de laborator',
            'perioada': {'de_la': None, 'pana_la': None},
            'perioada_cpt': M.get('perioada'),
            'generat': datetime.date.today().isoformat(),
        },
        'zone': ZONE_NUME, 'stadiu': stadiu,
        'investigatii': investigatii,
        'puncte_proiectate': [{'id': b, 'foraj': b, 'lat': v['lat'], 'lon': v['lon'], 'tip': 'Turbină (plan)', 'obs': 'plan 11.11.2025', 'adancime': '30 m'} for b, v in plan.items()],
        'cpt_neexecutate': [{'id': b, 'grup': b, 'turbina': b, 'lat': plan[b]['lat'], 'lon': plan[b]['lon']} for b in order if b not in exec_cpt_turb],
        'cpt': {k: M.get(k) for k in ('nr', 'nr_cptu', 'nr_turbine', 'ml', 'perioada', 'echipament', 'ipoteze', 'sbt', 'sbt_culori', 'dz')} | {'teste': light},
        'galerie_zona': {'titlu': 'Zona WT41 – WT42 (cea mai înaltă)', 'poze': zona_gal},
        'meteo': meteo,
    }
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, 'data.json'), 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(',', ':'))
    print('scris data.json', len(investigatii), 'foraje,', sum(len(i['lab']) for i in investigatii), 'probe lab,', len(light), 'CPT,',
          os.path.getsize(os.path.join(a.out, 'data.json')) // 1024, 'kB')
    for s in stadiu: print(' ', s)
    for i in investigatii:
        if len(i['nota_data']) > 120: print(' ', i['id'], i['nota_data'])


if __name__ == '__main__':
    main()
